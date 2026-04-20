"""
Fetches the top 25 most-viewed videos from a YouTube channel and saves the
results to an Excel file with columns:
  Rank | Title | Views | Link | Thumbnail URL | Thumbnail (embedded image)
"""

import io
import sys

import requests
import yt_dlp
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

CHANNEL_URL = "https://www.youtube.com/@NovaIVFFertility/videos"
OUTPUT_FILE = "NovaIVF_Top25_Videos.xlsx"
TOP_N = 25


def fetch_all_videos(channel_url: str) -> list[dict]:
    """Use yt-dlp to list all videos on the channel (flat extraction)."""
    ydl_opts = {
        "quiet": True,
        "extract_flat": "in_playlist",
        "skip_download": True,
        "playlistend": 500,  # cap at 500 to keep it reasonable
    }
    with yt_dlp.YoutubeDL(ydl_opts) as ydl:
        info = ydl.extract_info(channel_url, download=False)

    entries = info.get("entries", [])
    return entries


def fetch_video_details(video_id: str) -> dict:
    """Fetch full metadata (incl. view_count) for a single video."""
    url = f"https://www.youtube.com/watch?v={video_id}"
    ydl_opts = {
        "quiet": True,
        "skip_download": True,
        "no_warnings": True,
    }
    with yt_dlp.YoutubeDL(ydl_opts) as ydl:
        info = ydl.extract_info(url, download=False)
    return info


def download_thumbnail(url: str) -> bytes | None:
    """Download thumbnail bytes; return None on failure."""
    try:
        resp = requests.get(url, timeout=10)
        resp.raise_for_status()
        return resp.content
    except Exception:
        return None


def build_excel(videos: list[dict], output_path: str) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Top 25 Videos"

    # ── Header row ──────────────────────────────────────────────────────────
    headers = ["Rank", "Title", "Views", "Link", "Thumbnail URL", "Thumbnail"]
    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(color="FFFFFF", bold=True, size=11)
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    # Column widths (approx)
    col_widths = [6, 55, 14, 45, 55, 22]
    for i, w in enumerate(col_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ROW_HEIGHT = 80  # pixels-ish

    # ── Data rows ───────────────────────────────────────────────────────────
    for rank, video in enumerate(videos, 1):
        row = rank + 1
        ws.row_dimensions[row].height = ROW_HEIGHT

        title = video.get("title", "N/A")
        views = video.get("view_count", 0)
        video_id = video.get("id", "")
        link = f"https://www.youtube.com/watch?v={video_id}"
        thumbnail_url = video.get("thumbnail", "") or f"https://img.youtube.com/vi/{video_id}/hqdefault.jpg"

        ws.cell(row=row, column=1, value=rank).alignment = Alignment(horizontal="center", vertical="center")
        ws.cell(row=row, column=2, value=title).alignment = Alignment(wrap_text=True, vertical="center")
        views_cell = ws.cell(row=row, column=3, value=views)
        views_cell.alignment = Alignment(horizontal="right", vertical="center")
        views_cell.number_format = "#,##0"

        link_cell = ws.cell(row=row, column=4, value=link)
        link_cell.hyperlink = link
        link_cell.font = Font(color="0563C1", underline="single")
        link_cell.alignment = Alignment(vertical="center")

        ws.cell(row=row, column=5, value=thumbnail_url).alignment = Alignment(vertical="center", wrap_text=True)

        # Embed thumbnail image
        thumb_bytes = download_thumbnail(thumbnail_url)
        if thumb_bytes:
            try:
                img = XLImage(io.BytesIO(thumb_bytes))
                img.width = 120
                img.height = 68
                anchor_cell = f"F{row}"
                ws.add_image(img, anchor_cell)
            except Exception as exc:
                print(f"  [warn] Could not embed thumbnail for row {row}: {exc}")

        print(f"  [{rank:2d}] {title[:60]} — {views:,} views")

    # Freeze header row
    ws.freeze_panes = "A2"

    wb.save(output_path)
    print(f"\n✅  Saved: {output_path}")


def main() -> None:
    print("🔍  Fetching video list from channel…")
    entries = fetch_all_videos(CHANNEL_URL)
    if not entries:
        print("❌  No videos found. Exiting.")
        sys.exit(1)

    print(f"   Found {len(entries)} videos. Fetching view counts…")

    enriched = []
    for i, entry in enumerate(entries, 1):
        vid_id = entry.get("id") or entry.get("url", "").split("v=")[-1]
        if not vid_id:
            continue
        print(f"  ({i}/{len(entries)}) Fetching details for {vid_id}…", end="\r")
        try:
            details = fetch_video_details(vid_id)
            enriched.append(details)
        except Exception as exc:
            print(f"\n  [warn] Skipping {vid_id}: {exc}")

    if not enriched:
        print("❌  Could not retrieve any video details. Exiting.")
        sys.exit(1)

    # Sort by view count descending, take top N
    enriched.sort(key=lambda v: v.get("view_count") or 0, reverse=True)
    top25 = enriched[:TOP_N]

    print(f"\n📊  Building Excel file with top {len(top25)} videos…")
    build_excel(top25, OUTPUT_FILE)


if __name__ == "__main__":
    main()
